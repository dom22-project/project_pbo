"""
Script untuk create sample test file yang sudah diperbaiki
Jalankan: python create_fixed_test_file.py

Ini akan buat file Excel dengan struktur yang benar sehingga import berhasil
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

def create_test_file():
    """Create test Excel file dengan struktur yang benar"""
    
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    # Define styles
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    center_alignment = Alignment(horizontal='center', vertical='center')
    
    # Sheet 1: db table operasi
    ws1 = wb.create_sheet('db table operasi')
    headers1 = ['No', 'Fee Operator', 'Kelas', 'Harga Operator', 'Harga Anestesi']
    
    for col, header in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = center_alignment
    
    # FIXED DATA: Hanya data yang valid, tanpa kategori header
    sample_data1 = [
        [1, 'PERCUTANEOUS TRANSLUMINAL ANGIOPLASTY', 'ED', 18571500, 500000],
        [2, 'PERCUTANEOUS TRANSLUMINAL ANGIOPLASTY', 'OPD', 14285800, 250000],
        [3, 'PERCUTANEOUS TRANSLUMINAL ANGIOPLASTY', 'ODC', 14285800, 250000],
        [4, 'PERCUTANEOUS TRANSLUMINAL ANGIOPLASTY', 'BASIC', 14285800, 250000],
        [5, 'PERCUTANEOUS TRANSLUMINAL ANGIOPLASTY', 'STANDARD', 18571500, 300000],
        [6, 'PERCUTANEOUS TRANSLUMINAL ANGIOPLASTY', 'DELUXE', 20000000, 350000],
        [7, 'PERCUTANEOUS TRANSLUMINAL ANGIOPLASTY', 'VIP', 21428600, 400000],
        [8, 'PERCUTANEOUS TRANSLUMINAL ANGIOPLASTY', 'VVIP', 22857200, 450000],
        [9, 'PERCUTANEOUS TRANSLUMINAL ANGIOPLASTY', 'SUITE', 22857200, 450000],
        [10, 'PERCUTANEOUS TRANSLUMINAL ANGIOPLASTY', 'PRESIDENTIAL SUITE', 22857200, 500000],
        
        [11, 'CORONARY ANGIOGRAPHY', 'ED', 12000000, 300000],
        [12, 'CORONARY ANGIOGRAPHY', 'OPD', 10000000, 250000],
        [13, 'CORONARY ANGIOGRAPHY', 'ODC', 10000000, 250000],
        [14, 'CORONARY ANGIOGRAPHY', 'BASIC', 10000000, 250000],
        [15, 'CORONARY ANGIOGRAPHY', 'STANDARD', 12000000, 300000],
        [16, 'CORONARY ANGIOGRAPHY', 'DELUXE', 14000000, 350000],
        [17, 'CORONARY ANGIOGRAPHY', 'VIP', 15000000, 400000],
        [18, 'CORONARY ANGIOGRAPHY', 'VVIP', 16000000, 450000],
        [19, 'CORONARY ANGIOGRAPHY', 'SUITE', 16000000, 450000],
        [20, 'CORONARY ANGIOGRAPHY', 'PRESIDENTIAL SUITE', 16000000, 500000],
    ]
    
    for row_idx, row_data in enumerate(sample_data1, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=value)
            cell.border = border
            if col_idx > 1:
                cell.alignment = center_alignment
    
    # Set column widths
    ws1.column_dimensions['A'].width = 5
    ws1.column_dimensions['B'].width = 40
    ws1.column_dimensions['C'].width = 20
    ws1.column_dimensions['D'].width = 18
    ws1.column_dimensions['E'].width = 18
    
    # Sheet 2: db nama dokter
    ws2 = wb.create_sheet('db nama dokter')
    headers2 = ['No', 'Nama Dokter']
    
    for col, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = center_alignment
    
    sample_data2 = [
        [1, 'Dr. Budi Santoso, SpJP'],
        [2, 'Dr. Siti Nurhaliza, SpJP'],
        [3, 'Dr. Ahmad Wijaya, SpJP'],
        [4, 'Prof. Dr. Bambang Setianto, SpJP(K)'],
        [5, 'Dr. Hendra Wijaya, SpPD(K)'],
    ]
    
    for row_idx, row_data in enumerate(sample_data2, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=value)
            cell.border = border
            if col_idx > 1:
                cell.alignment = center_alignment
    
    ws2.column_dimensions['A'].width = 5
    ws2.column_dimensions['B'].width = 40
    
    # Save file
    filename = f'test_fixed_database_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    wb.save(filename)
    print(f"✓ File created: {filename}")
    print(f"  - Operasi: 20 data (2 operasi × 10 kelas)")
    print(f"  - Dokter: 5 data")
    print(f"\nFile ini SUDAH DIPERBAIKI dan siap untuk upload!")

if __name__ == '__main__':
    create_test_file()
