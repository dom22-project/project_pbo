"""Debug script untuk melihat struktur file Excel"""
import openpyxl
import os

excel_file = r"c:\Users\agung.daniel\Project PBO\app pbo\data\db pbo.xlsx"

if not os.path.exists(excel_file):
    print(f"File not found: {excel_file}")
    exit(1)

print(f"Opening: {excel_file}\n")

wb = openpyxl.load_workbook(excel_file)
print(f"Available sheets: {wb.sheetnames}\n")

# Check operasi sheet
sheet_name = "db table operasi"
if sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"{'='*80}")
    print(f"Sheet: {sheet_name}")
    print(f"Max row: {ws.max_row}, Max column: {ws.max_column}")
    print(f"{'='*80}\n")
    
    # Print first 5 rows
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row_idx > 5:
            break
        print(f"Row {row_idx}: {row}")
        if row_idx == 1:
            print(f"  -> Header dengan {len([c for c in row if c])} columns berisi data\n")
        else:
            print(f"  -> {len(row)} total columns, {len([c for c in row if c])} berisi data\n")

else:
    print(f"Sheet '{sheet_name}' tidak ditemukan!")
    print(f"Sheet yang tersedia: {wb.sheetnames}")

wb.close()
