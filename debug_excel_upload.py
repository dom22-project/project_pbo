"""
Script untuk debug dan analyze file Excel yang di-upload
Jalankan: python debug_excel_upload.py <path_file_excel>
"""

import sys
import openpyxl
from openpyxl.utils import get_column_letter

def debug_excel_file(file_path):
    """Analyze struktur dan isi file Excel"""
    
    print("=" * 80)
    print(f"DEBUG EXCEL FILE: {file_path}")
    print("=" * 80)
    
    try:
        wb = openpyxl.load_workbook(file_path)
        print(f"\n✓ File terbuka berhasil")
        print(f"  Total sheets: {len(wb.sheetnames)}")
        print(f"  Sheet names: {wb.sheetnames}\n")
        
        # Check setiap sheet
        for sheet_name in wb.sheetnames:
            print(f"\n{'='*80}")
            print(f"SHEET: {sheet_name}")
            print(f"{'='*80}")
            
            ws = wb[sheet_name]
            print(f"Rows: {ws.max_row}, Columns: {ws.max_column}")
            
            # Print first 10 rows
            print("\nFirst 10 rows:")
            print("-" * 80)
            
            for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                if row_idx > 10:
                    break
                
                # Truncate long values for display
                display_row = []
                for val in row:
                    if val is None:
                        display_row.append("[EMPTY]")
                    elif isinstance(val, str):
                        val_str = str(val)[:30]  # Truncate to 30 chars
                        if len(str(val)) > 30:
                            val_str += "..."
                        display_row.append(val_str)
                    else:
                        display_row.append(str(val)[:30])
                
                print(f"Row {row_idx:3d}: {display_row}")
            
            # Analyze data quality
            print(f"\n{'Data Quality Analysis:':^80}")
            print("-" * 80)
            
            empty_rows = 0
            rows_with_data = 0
            empty_columns_in_row = {}
            
            for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                if row_idx == 1:
                    continue  # Skip header
                
                if not row or not any(row):
                    empty_rows += 1
                    continue
                
                rows_with_data += 1
                
                # Count empty cells in this row
                empty_cells = sum(1 for cell in row if cell is None or cell == '')
                if empty_cells not in empty_columns_in_row:
                    empty_columns_in_row[empty_cells] = 0
                empty_columns_in_row[empty_cells] += 1
            
            print(f"Total rows (including header): {ws.max_row}")
            print(f"Total columns: {ws.max_column}")
            print(f"Rows with data: {rows_with_data}")
            print(f"Empty rows: {empty_rows}")
            print(f"Rows per empty cells pattern: {empty_columns_in_row}")
            
            # Sheet-specific analysis
            if "operasi" in sheet_name.lower():
                print(f"\n{'Analisis Operasi Sheet':^80}")
                print("-" * 80)
                check_operasi_sheet(ws)
            
            elif "dokter" in sheet_name.lower():
                print(f"\n{'Analisis Dokter Sheet':^80}")
                print("-" * 80)
                check_dokter_sheet(ws)
            
            elif "tindakan" in sheet_name.lower():
                print(f"\n{'Analisis Tindakan Sheet':^80}")
                print("-" * 80)
                check_tindakan_sheet(ws)
        
        wb.close()
        
    except Exception as e:
        print(f"✗ Error: {str(e)}")
        import traceback
        traceback.print_exc()

def check_operasi_sheet(ws):
    """Check operasi sheet format"""
    
    print("Expected format:")
    print("  Row 1: Headers - No, Fee Operator, Kelas, Harga Operator, Harga Anestesi")
    print("  Rows 2+: Data")
    
    # Check header
    header = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    print(f"\nActual header: {header}")
    
    expected_headers = ['No', 'Fee Operator', 'Kelas', 'Harga Operator', 'Harga Anestesi']
    print(f"Expected: {expected_headers}")
    
    # Check data rows
    valid_rows = 0
    rows_missing_fee_operator = 0
    rows_missing_kelas = 0
    
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row_idx == 1:
            continue
        
        if not row or not any(row):
            continue
        
        # Column indices (0-based)
        no = row[0] if len(row) > 0 else None
        fee_operator = row[1] if len(row) > 1 else None
        kelas = row[2] if len(row) > 2 else None
        harga_operator = row[3] if len(row) > 3 else None
        harga_anestesi = row[4] if len(row) > 4 else None
        
        if not fee_operator:
            rows_missing_fee_operator += 1
        elif not kelas:
            rows_missing_kelas += 1
        else:
            valid_rows += 1
    
    print(f"\nValid rows (dengan Fee Operator & Kelas): {valid_rows}")
    print(f"Rows missing Fee Operator: {rows_missing_fee_operator}")
    print(f"Rows missing Kelas: {rows_missing_kelas}")
    
    if valid_rows == 0:
        print("\n⚠️  MASALAH: Tidak ada row valid!")
        print("   Kemungkinan:")
        print("   1. Semua data di kolom B (Fee Operator) kosong")
        print("   2. Semua data di kolom C (Kelas) kosong")
        print("   3. Header tidak di row 1")

def check_dokter_sheet(ws):
    """Check dokter sheet format"""
    
    print("Expected format:")
    print("  Row 1: Headers - No, Nama Dokter")
    print("  Rows 2+: Data")
    
    # Check header
    header = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    print(f"\nActual header: {header}")
    
    # Check data rows
    valid_rows = 0
    rows_empty_name = 0
    
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row_idx == 1:
            continue
        
        if not row or not any(row):
            continue
        
        nama_dokter = row[1] if len(row) > 1 and row[1] else row[0]
        
        if nama_dokter:
            valid_rows += 1
        else:
            rows_empty_name += 1
    
    print(f"\nValid rows (dengan Nama Dokter): {valid_rows}")
    print(f"Rows empty name: {rows_empty_name}")

def check_tindakan_sheet(ws):
    """Check tindakan sheet format"""
    
    print("Expected format:")
    print("  Row 1: Headers - No, Nama Tindakan, Kelas, Kategory, Sales Item Type, Amount")
    print("  Rows 2+: Data")
    
    # Check header
    header = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    print(f"\nActual header: {header}")
    
    # Check data rows
    valid_rows = 0
    rows_missing_nama = 0
    rows_missing_kelas = 0
    
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row_idx == 1:
            continue
        
        if not row or not any(row):
            continue
        
        nama_tindakan = row[1] if len(row) > 1 else None
        kelas = row[2] if len(row) > 2 else None
        
        if not nama_tindakan:
            rows_missing_nama += 1
        elif not kelas:
            rows_missing_kelas += 1
        else:
            valid_rows += 1
    
    print(f"\nValid rows (dengan Nama & Kelas): {valid_rows}")
    print(f"Rows missing Nama Tindakan: {rows_missing_nama}")
    print(f"Rows missing Kelas: {rows_missing_kelas}")

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python debug_excel_upload.py <path_file_excel>")
        print("\nExample:")
        print("  python debug_excel_upload.py uploads/database.xlsx")
        sys.exit(1)
    
    file_path = sys.argv[1]
    debug_excel_file(file_path)
