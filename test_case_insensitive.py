#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Test case-insensitive sheet matching
"""

import openpyxl
from openpyxl.styles import Font, PatternFill
from pathlib import Path

def create_test_excel_uppercase():
    """Create test Excel file dengan uppercase sheet names"""
    wb = openpyxl.Workbook()
    
    # Create db table operasi sheet dengan UPPERCASE
    ws_operasi = wb.active
    ws_operasi.title = 'DB TABLE OPERASI'  # UPPERCASE!
    
    # Add headers
    headers_operasi = ['No', 'Fee Operator', 'Kelas', 'Harga Operator', 'Harga Anestesi']
    for col_idx, header in enumerate(headers_operasi, 1):
        cell = ws_operasi.cell(row=1, column=col_idx)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        cell.font = Font(bold=True, color='FFFFFF')
    
    # Add sample data
    ws_operasi['A2'] = 1
    ws_operasi['B2'] = 'Op1'
    ws_operasi['C2'] = 'GENERAL'
    ws_operasi['D2'] = 1000000
    ws_operasi['E2'] = 500000
    
    # Create db nama dokter sheet dengan UPPERCASE
    ws_dokter = wb.create_sheet('DB NAMA DOKTER')  # UPPERCASE!
    
    # Add headers
    headers_dokter = ['No', 'Nama Dokter']
    for col_idx, header in enumerate(headers_dokter, 1):
        cell = ws_dokter.cell(row=1, column=col_idx)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')
        cell.font = Font(bold=True, color='FFFFFF')
    
    # Add doctor data
    doctors = [
        'Dr. Johndoe',
        'Dr. Janedoe',
        'Dr. Smith',
        'Dr. Johnson'
    ]
    
    for row_idx, doctor_name in enumerate(doctors, 2):
        ws_dokter.cell(row=row_idx, column=1).value = row_idx - 1
        ws_dokter.cell(row=row_idx, column=2).value = doctor_name
    
    # Save file
    test_file = Path('test_doctors_uppercase.xlsx')
    wb.save(str(test_file))
    print(f"[TEST] Test file created: {test_file}")
    print(f"[TEST] Sheet names: {wb.sheetnames}")
    
    wb.close()
    return str(test_file)

if __name__ == '__main__':
    filepath = create_test_excel_uppercase()
    
    # Now test import
    import os
    import sys
    sys.path.insert(0, str(Path(__file__).parent))
    os.environ['FLASK_ENV'] = 'development'
    
    from app import app, db_helper, import_excel_to_database
    
    print("\n=== Testing Doctor Import with UPPERCASE sheets ===\n")
    
    with app.app_context():
        db_helper.delete_all_doctors()
        print(f"[TEST] Current doctors count: {db_helper.count_doctors()}")
        
        print(f"\n[TEST] Starting import from: {filepath}")
        try:
            stats = import_excel_to_database(filepath, db_helper)
            
            print("\n[TEST] Import Statistics:")
            print(f"  - Doctors imported: {stats.get('doctors_imported', 0)}")
            print(f"  - Doctors duplicates: {stats.get('doctors_duplicates', 0)}")
            print(f"  - Doctors skipped: {stats.get('doctors_skipped', 0)}")
            print(f"  - Warnings: {stats.get('warnings', [])}")
            
            # Verify doctors in database
            print(f"\n[TEST] Verifying doctors in database...")
            all_doctors = db_helper.get_all_doctors()
            print(f"[TEST] Total doctors in DB: {len(all_doctors)}")
            
            if all_doctors:
                print("\n[TEST] Doctors in database:")
                for doctor in all_doctors:
                    print(f"  - {doctor['nama_dokter']}")
            else:
                print("[ERROR] No doctors found in database after import!")
                
        except Exception as e:
            print(f"[ERROR] Import failed: {str(e)}")
            import traceback
            traceback.print_exc()
