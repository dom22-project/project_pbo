"""Test direct import function"""
import sys
sys.path.insert(0, r"c:\Users\agung.daniel\Project PBO\app pbo")

from models import Database
import openpyxl

excel_file = r"c:\Users\agung.daniel\Project PBO\app pbo\data\db pbo.xlsx"

print(f"Opening Excel: {excel_file}")
wb = openpyxl.load_workbook(excel_file)

sheet_name = "db table operasi"
ws = wb[sheet_name]

print(f"\n=== CHECKING FIRST 10 DATA ROWS ===\n")

header_row_idx = None
for idx, row in enumerate(ws.iter_rows(values_only=True), 1):
    if idx <= 5:
        print(f"Row {idx}: {row}")
    if row and any(row) and header_row_idx is None:
        header_row_idx = idx
        print(f"\n*** HEADER FOUND AT ROW {idx} ***\n")

print(f"\nProcessing data starting from row {header_row_idx + 1}...\n")

count = 0
for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
    if row_idx <= header_row_idx:
        continue
    if not row or not any(row):
        continue
    
    count += 1
    if count <= 10:
        no = row[0]
        fee_operator = row[1] if len(row) > 1 else None
        kelas = row[2] if len(row) > 2 else None
        harga_operator = row[3] if len(row) > 3 else None
        harga_anestesi = row[4] if len(row) > 4 else None
        
        print(f"Data Row {count}:")
        print(f"  no={no}, fee_operator={fee_operator}, kelas={kelas}")
        print(f"  harga_operator={harga_operator}, harga_anestesi={harga_anestesi}")
        print(f"  fee_operator is None: {fee_operator is None}")
        print(f"  kelas is None: {kelas is None}")
        print()

print(f"Total data rows (non-empty): {count}")
wb.close()
