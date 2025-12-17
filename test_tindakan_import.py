"""
Test script untuk memverifikasi import data tindakan items
"""

from models import Database
import os

def test_database_table():
    """Test 1: Verifikasi tabel tindakan_items ada di database"""
    print("=" * 60)
    print("TEST 1: Verifikasi Tabel Database")
    print("=" * 60)
    
    db = Database()
    conn = db.get_connection()
    cursor = conn.cursor()
    
    # Check if table exists
    cursor.execute("""
        SELECT name FROM sqlite_master 
        WHERE type='table' AND name='tindakan_items'
    """)
    
    result = cursor.fetchone()
    conn.close()
    
    if result:
        print("✓ Tabel 'tindakan_items' berhasil dibuat")
        return True
    else:
        print("✗ Tabel 'tindakan_items' tidak ditemukan")
        return False

def test_count_tindakan():
    """Test 2: Hitung jumlah data tindakan items"""
    print("\n" + "=" * 60)
    print("TEST 2: Hitung Data Tindakan Items")
    print("=" * 60)
    
    db = Database()
    count = db.count_tindakan_items()
    
    print(f"Total tindakan items di database: {count}")
    
    if count > 0:
        print(f"✓ Berhasil menemukan {count} tindakan items")
        return True
    else:
        print("⚠ Tidak ada data tindakan items (mungkin sheet tidak ada atau kosong)")
        return False

def test_get_tindakan_items():
    """Test 3: Ambil dan tampilkan sample data tindakan items"""
    print("\n" + "=" * 60)
    print("TEST 3: Ambil Sample Data Tindakan Items")
    print("=" * 60)
    
    db = Database()
    items = db.get_all_tindakan_items()
    
    if items:
        print(f"✓ Berhasil mengambil {len(items)} tindakan items")
        print("\nSample data (5 pertama):")
        print("-" * 60)
        
        for i, item in enumerate(items[:5], 1):
            print(f"\n{i}. {item['nama_tindakan']}")
            print(f"   Kelas: {item['kelas']}")
            print(f"   Kategory: {item['kategory']}")
            print(f"   Sales Item Type: {item['sales_item_type']}")
            print(f"   Amount: Rp {item['amount']:,.0f}")
        
        return True
    else:
        print("✗ Tidak ada data tindakan items")
        return False

def test_add_tindakan_item():
    """Test 4: Test fungsi add_tindakan_item"""
    print("\n" + "=" * 60)
    print("TEST 4: Test Fungsi Add Tindakan Item")
    print("=" * 60)
    
    db = Database()
    
    # Get count before
    count_before = db.count_tindakan_items()
    
    # Add test item
    test_item = {
        'nama_tindakan': 'TEST TINDAKAN',
        'kelas': 'TEST CLASS',
        'kategory': 'TEST CATEGORY',
        'sales_item_type': 'TEST TYPE',
        'amount': 100000
    }
    
    try:
        item_id = db.add_tindakan_item(
            nama_tindakan=test_item['nama_tindakan'],
            kelas=test_item['kelas'],
            kategory=test_item['kategory'],
            sales_item_type=test_item['sales_item_type'],
            amount=test_item['amount']
        )
        
        # Get count after
        count_after = db.count_tindakan_items()
        
        if count_after == count_before + 1:
            print(f"✓ Berhasil menambahkan tindakan item dengan ID: {item_id}")
            print(f"  Jumlah sebelum: {count_before}")
            print(f"  Jumlah sesudah: {count_after}")
            
            # Clean up - delete test item
            conn = db.get_connection()
            cursor = conn.cursor()
            cursor.execute("DELETE FROM tindakan_items WHERE id = ?", (item_id,))
            conn.commit()
            conn.close()
            print("  Test item berhasil dihapus (cleanup)")
            
            return True
        else:
            print("✗ Gagal menambahkan tindakan item")
            return False
            
    except Exception as e:
        print(f"✗ Error saat menambahkan tindakan item: {str(e)}")
        return False

def test_database_structure():
    """Test 5: Verifikasi struktur tabel tindakan_items"""
    print("\n" + "=" * 60)
    print("TEST 5: Verifikasi Struktur Tabel")
    print("=" * 60)
    
    db = Database()
    conn = db.get_connection()
    cursor = conn.cursor()
    
    cursor.execute("PRAGMA table_info(tindakan_items)")
    columns = cursor.fetchall()
    conn.close()
    
    expected_columns = ['id', 'nama_tindakan', 'kelas', 'kategory', 'sales_item_type', 'amount', 'created_at']
    actual_columns = [col[1] for col in columns]
    
    print("Kolom yang ditemukan:")
    for col in columns:
        print(f"  - {col[1]} ({col[2]})")
    
    missing = set(expected_columns) - set(actual_columns)
    if not missing:
        print("\n✓ Semua kolom yang diharapkan ada")
        return True
    else:
        print(f"\n✗ Kolom yang hilang: {missing}")
        return False

def test_excel_file():
    """Test 6: Verifikasi file Excel dan sheet"""
    print("\n" + "=" * 60)
    print("TEST 6: Verifikasi File Excel")
    print("=" * 60)
    
    excel_path = 'data/db pbo.xlsx'
    
    if not os.path.exists(excel_path):
        print(f"✗ File Excel tidak ditemukan: {excel_path}")
        return False
    
    print(f"✓ File Excel ditemukan: {excel_path}")
    
    try:
        import openpyxl
        wb = openpyxl.load_workbook(excel_path)
        
        print(f"\nSheet yang tersedia:")
        for sheet in wb.sheetnames:
            print(f"  - {sheet}")
        
        if 'db nama tindakan' in wb.sheetnames:
            print("\n✓ Sheet 'db nama tindakan' ditemukan")
            
            ws = wb['db nama tindakan']
            row_count = ws.max_row - 1  # -1 untuk header
            print(f"  Jumlah baris data (tidak termasuk header): {row_count}")
            
            wb.close()
            return True
        else:
            print("\n✗ Sheet 'db nama tindakan' tidak ditemukan")
            wb.close()
            return False
            
    except Exception as e:
        print(f"✗ Error membaca file Excel: {str(e)}")
        return False

def run_all_tests():
    """Jalankan semua test"""
    print("\n" + "=" * 60)
    print("TESTING IMPORT DATA TINDAKAN ITEMS")
    print("=" * 60)
    
    tests = [
        ("Verifikasi Tabel Database", test_database_table),
        ("Verifikasi Struktur Tabel", test_database_structure),
        ("Verifikasi File Excel", test_excel_file),
        ("Hitung Data Tindakan", test_count_tindakan),
        ("Ambil Sample Data", test_get_tindakan_items),
        ("Test Add Function", test_add_tindakan_item),
    ]
    
    results = []
    
    for test_name, test_func in tests:
        try:
            result = test_func()
            results.append((test_name, result))
        except Exception as e:
            print(f"\n✗ Error pada test '{test_name}': {str(e)}")
            results.append((test_name, False))
    
    # Summary
    print("\n" + "=" * 60)
    print("RINGKASAN TEST")
    print("=" * 60)
    
    passed = sum(1 for _, result in results if result)
    total = len(results)
    
    for test_name, result in results:
        status = "✓ PASS" if result else "✗ FAIL"
        print(f"{status} - {test_name}")
    
    print("\n" + "=" * 60)
    print(f"Total: {passed}/{total} tests passed")
    print("=" * 60)
    
    if passed == total:
        print("\n🎉 Semua test berhasil!")
    else:
        print(f"\n⚠ {total - passed} test gagal")
    
    return passed == total

if __name__ == "__main__":
    success = run_all_tests()
    exit(0 if success else 1)
